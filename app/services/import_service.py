"""Source-document parsing and asynchronous novel/style import workflows."""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
import json
import logging
from pathlib import Path
import re
import threading
from zipfile import BadZipFile, ZipFile

from sqlalchemy import select
from sqlalchemy.orm import Session

from app import prompts
from app.config import get_settings
from app.db import session_scope
from app.models import Chapter, Novel, NovelSourceReference, SourceDocument
from app.providers.registry import get_registry

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".txt", ".md", ".markdown", ".docx"}

# These semaphores deliberately reject excess work instead of creating an
# unbounded number of daemon threads that could exhaust the model proxy or CPU.
_parse_slots = threading.BoundedSemaphore(
    max(1, get_settings().source_parse_max_concurrency)
)
_novel_import_slots = threading.BoundedSemaphore(
    max(1, get_settings().import_worker_max_concurrency)
)
_style_analysis_slots = threading.BoundedSemaphore(
    max(1, get_settings().style_analysis_max_concurrency)
)

_CHINESE_CHAPTER_HEADING = re.compile(
    r"^\s*(第\s*[0-9０-９零〇一二三四五六七八九十百千万两]+\s*(?:章|节|回|卷).{0,100})\s*$"
)
_MARKDOWN_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.{1,120}?)\s*#*\s*$")


class SourceImportError(ValueError):
    """Raised for an invalid or unsupported imported source file."""


class SourceImportBusyError(SourceImportError):
    """Raised when bounded source-processing capacity is currently full."""


@dataclass(frozen=True)
class ParsedChapter:
    title: str
    content: str


def _safe_title(value: str, fallback: str) -> str:
    text = " ".join((value or "").strip().split())
    return (text or fallback)[:255]


def _source_stem(filename: str) -> str:
    stem = Path(filename or "导入小说").stem.strip()
    return _safe_title(stem, "导入小说")


def _normalize_text(value: str) -> str:
    text = value.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    # Preserve paragraph boundaries but prevent an accidental huge block of blank lines.
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _decode_text(payload: bytes) -> str:
    if payload.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return payload.decode("utf-16")
        except UnicodeDecodeError as exc:
            raise SourceImportError("UTF-16 文本编码无效。") from exc
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise SourceImportError("无法识别文本编码，请保存为 UTF-8、UTF-16 或 GB18030 后重试。")


def _parse_upload_bytes(filename: str, payload: bytes) -> tuple[str, str]:
    """Return normalized text and the source extension after validation."""
    extension = Path(filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        allowed = "、".join(sorted(ALLOWED_EXTENSIONS))
        raise SourceImportError(f"仅支持 {allowed} 格式的小说文件。")
    if not payload:
        raise SourceImportError("上传的文件为空。")

    settings = get_settings()
    if extension == ".docx":
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency deployment failure
            raise SourceImportError("服务器缺少 DOCX 解析依赖。") from exc
        try:
            # A DOCX is a ZIP archive.  Limiting its compressed upload size is
            # not enough: a tiny archive can otherwise expand to gigabytes.
            with ZipFile(BytesIO(payload)) as archive:
                members = archive.infolist()
                if len(members) > 2000:
                    raise SourceImportError("DOCX 文件包含过多内部文件，无法安全导入。")
                expanded_size = sum(member.file_size for member in members)
                if expanded_size > settings.source_docx_max_uncompressed_bytes:
                    raise SourceImportError("DOCX 解压后的内容超过服务器允许的大小。")
                if any(
                    member.file_size > 2 * 1024 * 1024
                    and member.compress_size > 0
                    and member.file_size / member.compress_size > 300
                    for member in members
                ):
                    raise SourceImportError("DOCX 压缩比例异常，无法安全导入。")
        except SourceImportError:
            raise
        except BadZipFile as exc:
            raise SourceImportError("DOCX 文件格式无效。") from exc
        try:
            document = Document(BytesIO(payload))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:
            raise SourceImportError("无法读取该 DOCX 文件。") from exc
    else:
        text = _decode_text(payload)

    text = _normalize_text(text)
    if not text:
        raise SourceImportError("文件中没有可导入的正文内容。")
    if len(text) > settings.source_max_extracted_chars:
        raise SourceImportError(
            f"提取出的正文超过 {settings.source_max_extracted_chars:,} 个字符上限。"
        )
    return text, extension


def parse_upload_bytes(filename: str, payload: bytes) -> tuple[str, str]:
    """Parse source text with bounded CPU/archive-expansion capacity."""
    if not _parse_slots.acquire(blocking=False):
        raise SourceImportBusyError("服务器正在处理较多导入文件，请稍后再试。")
    try:
        return _parse_upload_bytes(filename, payload)
    finally:
        _parse_slots.release()


def split_chapters(text: str) -> list[ParsedChapter]:
    """Split common Chinese/Markdown headings, falling back to one chapter."""
    current_title = ""
    current_lines: list[str] = []
    parsed: list[ParsedChapter] = []
    saw_heading = False

    def flush() -> None:
        content = _normalize_text("\n".join(current_lines))
        if content:
            parsed.append(
                ParsedChapter(
                    title=_safe_title(current_title, f"第 {len(parsed) + 1} 章"),
                    content=content,
                )
            )

    for line in text.splitlines():
        match = _CHINESE_CHAPTER_HEADING.match(line) or _MARKDOWN_HEADING.match(line)
        if match:
            saw_heading = True
            flush()
            current_lines = []
            current_title = match.group(1).strip()
        else:
            current_lines.append(line)
    flush()

    if not saw_heading or not parsed:
        content = _normalize_text(text)
        return [ParsedChapter(title="导入内容", content=content)] if content else []
    return parsed


def _chapter_summary(content: str) -> str:
    compact = re.sub(r"\s+", " ", content).strip()
    if len(compact) <= 420:
        return compact
    return f"{compact[:360]}……{compact[-60:]}"


def _sample_for_analysis(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    third = max(1, max_chars // 3)
    middle = max(0, len(text) // 2 - third // 2)
    return "\n\n……（中间内容已省略）……\n\n".join(
        (text[:third], text[middle : middle + third], text[-third:])
    )


def _sample_indexes(length: int, max_items: int) -> list[int]:
    if length <= max_items:
        return list(range(length))
    return sorted(
        {
            round(position * (length - 1) / (max_items - 1))
            for position in range(max_items)
        }
    )


def _heuristic_story_profile(chapters: list[ParsedChapter]) -> str:
    """Build a bounded, cross-book continuity map without requiring an LLM.

    Imported works must remain usable when the model proxy is unavailable.  A
    sample spread across the whole book preserves middle-plot developments far
    better than only retaining the first and final chapters.
    """
    if not chapters:
        return ""
    lines = [f"导入小说共 {len(chapters)} 章。以下为跨章节剧情连续性档案："]
    used = len(lines[0])
    for index in _sample_indexes(len(chapters), max_items=28):
        chapter = chapters[index]
        summary = _chapter_summary(chapter.content)[:190]
        line = f"{index + 1}. {chapter.title}：{summary}"
        if used + len(line) > 6000:
            lines.append("（其余章节请通过章节目录与最近路径摘要衔接。）")
            break
        lines.append(line)
        used += len(line)
    return "\n".join(lines)


def _import_outline(chapters: list[ParsedChapter]) -> str:
    lines: list[str] = []
    used = 0
    for index in _sample_indexes(len(chapters), max_items=80):
        line = f"{index + 1}. {chapters[index].title}"
        if used + len(line) > 5000:
            break
        lines.append(line)
        used += len(line)
    return "\n".join(lines)


def _load_options(document: SourceDocument) -> dict[str, str]:
    try:
        raw = json.loads(document.options_json or "{}")
    except json.JSONDecodeError:
        raw = {}
    return {str(key): str(value) for key, value in raw.items() if value is not None}


def _set_failed(document_id: int, message: str) -> None:
    with session_scope() as db:
        document = db.get(SourceDocument, document_id)
        if document is not None:
            document.status = "failed"
            document.error = message[:4000]


def _start_worker(target, *args) -> None:
    threading.Thread(target=target, args=args, daemon=True).start()


def _run_with_reserved_slot(slot: threading.BoundedSemaphore, target, *args) -> None:
    try:
        target(*args)
    finally:
        slot.release()


def create_novel_import(
    db: Session,
    *,
    owner_id: int,
    filename: str,
    content_type: str | None,
    payload: bytes,
    title: str | None,
    genre: str | None,
    style: str | None,
    provider: str | None,
    model: str | None,
) -> SourceDocument:
    settings = get_settings()
    if len(payload) > settings.upload_max_bytes:
        raise SourceImportError(
            f"文件超过上传上限（{settings.upload_max_bytes // (1024 * 1024)} MiB）。"
        )
    text, _extension = parse_upload_bytes(filename, payload)
    if not _novel_import_slots.acquire(blocking=False):
        raise SourceImportBusyError("服务器正在导入较多小说，请稍后再试。")
    options = {
        "title": _safe_title(title or "", _source_stem(filename)),
        "genre": (genre or "").strip()[:255],
        "style": (style or "").strip()[:255],
        "provider": (provider or "").strip(),
        "model": (model or "").strip(),
    }
    try:
        document = SourceDocument(
            owner_id=owner_id,
            kind="novel_import",
            title=options["title"],
            original_filename=(filename or "导入小说")[:512],
            mime_type=(content_type or "")[:255],
            size_bytes=len(payload),
            content_sha256=sha256(payload).hexdigest(),
            content=text,
            options_json=json.dumps(options, ensure_ascii=False),
            status="processing",
        )
        db.add(document)
        db.commit()
        db.refresh(document)
        _start_worker(_run_with_reserved_slot, _novel_import_slots, _process_novel_import, document.id)
        return document
    except Exception:
        _novel_import_slots.release()
        raise


def _process_novel_import(document_id: int) -> None:
    try:
        with session_scope() as db:
            document = db.get(SourceDocument, document_id)
            if document is None:
                return
            chapters = split_chapters(document.content)
            if not chapters:
                raise SourceImportError("未能从文件中提取到章节内容。")
            if len(chapters) > get_settings().source_max_chapters:
                raise SourceImportError(
                    f"识别出的章节超过 {get_settings().source_max_chapters:,} 章上限。"
                )
            options = _load_options(document)
            story_profile = _heuristic_story_profile(chapters)
            novel = Novel(
                owner_id=document.owner_id,
                title=_safe_title(options.get("title", ""), document.title or "导入小说"),
                genre=options.get("genre", ""),
                premise=story_profile,
                settings="",
                world_setting="",
                style=options.get("style", ""),
                outline=_import_outline(chapters),
                provider=options.get("provider") or None,
                model=options.get("model") or None,
                is_generating=False,
            )
            db.add(novel)
            db.flush()

            parent_id: int | None = None
            for index, parsed in enumerate(chapters, start=1):
                chapter = Chapter(
                    novel_id=novel.id,
                    parent_id=parent_id,
                    index=index,
                    title=parsed.title,
                    content=parsed.content,
                    summary=_chapter_summary(parsed.content),
                    plot_directive="由导入小说内容创建",
                    is_primary=parent_id is not None,
                    is_ending=False,
                )
                db.add(chapter)
                db.flush()
                parent_id = chapter.id

            db.add(
                NovelSourceReference(
                    novel_id=novel.id,
                    source_document_id=document.id,
                    usage="continuation",
                )
            )
            document.story_profile = story_profile
            document.result_novel_id = novel.id
            document.status = "ready"
            document.error = None
        logger.info("Imported source document %s into a novel", document_id)
    except Exception as exc:
        logger.exception("Novel import failed for source document %s", document_id)
        _set_failed(document_id, str(exc))


def create_style_reference(
    db: Session,
    *,
    owner_id: int,
    filename: str,
    content_type: str | None,
    payload: bytes,
    title: str | None,
) -> SourceDocument:
    settings = get_settings()
    if len(payload) > settings.upload_max_bytes:
        raise SourceImportError(
            f"文件超过上传上限（{settings.upload_max_bytes // (1024 * 1024)} MiB）。"
        )
    text, _extension = parse_upload_bytes(filename, payload)
    if not _style_analysis_slots.acquire(blocking=False):
        raise SourceImportBusyError("文风分析队列已满，请稍后再试。")
    try:
        document = SourceDocument(
            owner_id=owner_id,
            kind="style_reference",
            title=_safe_title(title or "", _source_stem(filename)),
            original_filename=(filename or "文风参考")[:512],
            mime_type=(content_type or "")[:255],
            size_bytes=len(payload),
            content_sha256=sha256(payload).hexdigest(),
            content=text,
            status="processing",
        )
        db.add(document)
        db.commit()
        db.refresh(document)
        _start_worker(
            _run_with_reserved_slot,
            _style_analysis_slots,
            _process_style_reference,
            document.id,
        )
        return document
    except Exception:
        _style_analysis_slots.release()
        raise


def _process_style_reference(document_id: int) -> None:
    try:
        with session_scope() as db:
            document = db.get(SourceDocument, document_id)
            if document is None:
                return
            sample = _sample_for_analysis(
                document.content, get_settings().source_analysis_max_chars
            )
            title = document.title

        provider = get_registry().get(None)
        profile = provider.generate(
            prompts.build_style_profile_prompt(sample, title),
            model=None,
            system=prompts.SYSTEM_PROMPT,
            temperature=0.25,
            timeout=get_settings().request_timeout,
        ).strip()
        if not profile:
            raise SourceImportError("模型未返回可用的文风档案。")

        with session_scope() as db:
            document = db.get(SourceDocument, document_id)
            if document is None:
                return
            document.style_profile = profile[:8000]
            document.status = "ready"
            document.error = None
        logger.info("Analyzed style reference %s", document_id)
    except Exception as exc:
        logger.exception("Style-reference analysis failed for source document %s", document_id)
        _set_failed(document_id, str(exc))


def list_style_references(db: Session, owner_id: int | None = None) -> list[SourceDocument]:
    query = (
        select(SourceDocument)
        .where(SourceDocument.kind == "style_reference")
        .order_by(SourceDocument.updated_at.desc())
    )
    if owner_id is not None:
        query = query.where(SourceDocument.owner_id == owner_id)
    return list(db.scalars(query))


def delete_source_document(db: Session, document: SourceDocument) -> None:
    """Erase stored source prose and its reusable reference links."""
    db.delete(document)
    db.commit()
